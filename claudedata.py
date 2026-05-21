# 04-XX 升級版：支援 5 大格式 (CSV, PDF, DOCX, JSON, TXT)
# 核心優化：Nomic 前綴對齊、修正欄位錯置、結構化語義重建
import os
import sys
import hashlib
import shutil
import json
import re
from glob import glob
from typing import List, Tuple, Dict

import pandas as pd

try:
    import docx
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False
    print("⚠️ 尚未安裝 python-docx，若需上傳 Word 檔案請先執行: pip install python-docx")

# 解決 Windows 權限與 Symlink 報錯
os.environ["HF_HUB_DISABLE_SYMLINKS"] = "1"
os.environ["HF_HUB_DISABLE_SYMLINKS_WARNING"] = "1"
sys.stdout.reconfigure(encoding="utf-8")

# Docling 與 LangChain 核心套件
from docling.document_converter import DocumentConverter, PdfFormatOption
from docling.datamodel.pipeline_options import PdfPipelineOptions, TableFormerMode
from docling.datamodel.base_models import InputFormat
from langchain_text_splitters import MarkdownHeaderTextSplitter, RecursiveCharacterTextSplitter
from langchain_ollama import OllamaEmbeddings, OllamaLLM
from langchain_chroma import Chroma
from langchain_core.documents import Document

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data")
DB_DIR = os.path.join(BASE_DIR, "chroma_db_unified")
EMBED_MODEL = "nomic-embed-text"
TAGGER_MODEL = "llama3.2" 

def get_advanced_pdf_converter() -> DocumentConverter:
    print("[系統配置] 正在初始化高階 OCR 與表格辨識管線...")
    pipeline_options = PdfPipelineOptions()
    pipeline_options.do_ocr = True 
    pipeline_options.do_table_structure = True 
    pipeline_options.table_structure_options.mode = TableFormerMode.ACCURATE 
    
    converter = DocumentConverter(
        allowed_formats=[InputFormat.PDF],
        format_options={
            InputFormat.PDF: PdfFormatOption(pipeline_options=pipeline_options)
        }
    )
    return converter

# ==========================================================
# 🌟 LLM 自動打標器 (修正 JSON 欄位對齊)
# ==========================================================
class LLM_Tagger:
    def __init__(self, model_name=TAGGER_MODEL):
        self.llm = OllamaLLM(model=model_name, temperature=0)
        self.prompt_template = """
        你是一個工業設備維修資料庫的管理員。請閱讀以下文本，並萃取關鍵資訊。
        
        【要求】
        1. 僅輸出 JSON 格式，絕對不要加上 Markdown 標記 (如 ```json) 或任何其他文字。
        2. JSON 必須包含以下三個欄位，鍵名必須完全精確：
           - "ai_equipment": <字串> (提及的設備或組件名稱。若無，填寫 "未知")
           - "ai_error_code": <字串> (提及的任何故障代碼或警報編號，多個用逗號分隔。若無，填寫 "無")
           - "ai_category": <字串> (從以下選擇最適合的一項："故障排除", "安全警告", "操作步驟", "規格參數", "保養維護", "其他")
           
        【文本內容】
        {text}
        """

    def tag(self, text: str) -> dict:
        default_tags = {"ai_equipment": "未知", "ai_error_code": "無", "ai_category": "其他"}
        # 移除 Nomic 訓練前綴後再送給 LLM 打標，避免干擾
        clean_text = text.replace("search_document: ", "").strip()

        if len(clean_text) < 20:
            return default_tags
        try:
            prompt = self.prompt_template.format(text=clean_text[:600])
            response = self.llm.invoke(prompt)
            match = re.search(r'\{[\s\S]*?\}', response, re.DOTALL)
            if match: 
                raw_json = json.loads(match.group(0))
                return {
                    "ai_equipment": str(raw_json.get("ai_equipment", "未知")),
                    "ai_error_code": str(raw_json.get("ai_error_code", "無")),
                    "ai_category": str(raw_json.get("ai_category", "其他"))
                }
        except Exception as e: 
            print(f"⚠️ 語意標籤提取或 JSON 解析失敗: {e}。將採用預設安全標籤。")
        return default_tags

tagger = LLM_Tagger()

def generate_doc_id(source: str, index: int) -> str:
    raw_id = f"{source}_{index}"
    return hashlib.md5(raw_id.encode("utf-8")).hexdigest()

# ==========================================================
# 🌟 各式檔案解析核心邏輯 (注入 Nomic 文件前綴)
# ==========================================================

def _read_csv_with_fallback(file_path: str) -> pd.DataFrame:
    encodings = ["utf-8-sig", "utf-8", "cp950", "big5"]
    last_err = None
    for enc in encodings:
        try:
            df = pd.read_csv(file_path, encoding=enc)
            df.columns = df.columns.str.strip()
            if '設備' in df.columns: df['設備'] = df['設備'].ffill()
            elif '組件' in df.columns: df['組件'] = df['組件'].ffill()
            return df.fillna("無")
        except Exception as e: last_err = e
    raise last_err

def process_single_csv(file_path: str) -> Tuple[list, list]:
    docs, ids = [], []
    try:
        df = _read_csv_with_fallback(file_path)
        source_name = os.path.basename(file_path)
        columns = list(df.columns)

        for i, row in df.iterrows():
            device = str(row.get('設備', row.get('組件', '未知'))).strip()
            code   = str(row.get('故障代碼', row.get('代碼', '無'))).strip()
            issue  = str(row.get('常見異常原因', row.get('異常問題', '未知'))).strip()
            check  = str(row.get('檢查方式', '無')).strip()
            sol    = str(row.get('解決方法', '無')).strip()
            loto   = str(row.get('安全警告與隔離(LOTO)', '無')).strip()
            note   = str(row.get('專家實務備註', '無')).strip()

            # 1. 結構化重組文本，大幅提升語義向量的捕捉能力
            # 2. 強制加上 nomic 必備的文件端前綴
            semantic_text = (
                f"search_document: 太陽能運維手冊 - 設備組件: {device}\n"
                f"故障代碼: {code}\n"
                f"異常現象與原因: {issue}\n"
                f"安全警告與隔離安全防護 (LOTO): {loto}\n"
                f"標準檢查作業流程 SOP: {check}\n"
                f"最終維修處置解決方法: {sol}\n"
                f"資深專家實務備註保養重點: {note}"
            )
            
            row_dict = {col: str(row[col]) for col in columns}
            llm_tags = tagger.tag(semantic_text)

            metadata = {
                "source": source_name, "type": "csv_log",
                "row_index": int(i), "col_count": len(columns),
                "columns_json": json.dumps(columns, ensure_ascii=False),
                "row_json": json.dumps(row_dict, ensure_ascii=False),
                "ai_equipment": llm_tags.get("ai_equipment", device),
                "ai_error_code": code if code != "無" else llm_tags.get("ai_error_code", "無"),
                "ai_category":  llm_tags.get("ai_category", "故障排除")
            }
            docs.append(Document(page_content=semantic_text, metadata=metadata))
            ids.append(generate_doc_id(source_name, i))
    except Exception as e:
        print(f"讀取 CSV 失敗 ({file_path}): {e}")
    return docs, ids

def process_single_pdf(file_path: str, converter=None) -> Tuple[list, list]:
    docs, ids = [], []
    try:
        if converter is None: converter = get_advanced_pdf_converter()
            
        md_splitter = MarkdownHeaderTextSplitter(headers_to_split_on=[("#", "H1"), ("##", "H2"), ("###", "H3")], strip_headers=False)
        text_splitter = RecursiveCharacterTextSplitter(chunk_size=1200, chunk_overlap=300)

        source_name = os.path.basename(file_path)
        print(f"  -> 正在解析: {source_name}")
        result  = converter.convert(file_path)
        doc_obj = result.document

        page_lookup: dict = {}
        try:
            for text_item in doc_obj.texts:
                if hasattr(text_item, "prov") and text_item.prov:
                    page_lookup[(text_item.text or "")[:30]] = text_item.prov[0].page_no
        except: pass 

        header_splits = md_splitter.split_text(doc_obj.export_to_markdown())
        final_chunks  = text_splitter.split_documents(header_splits)

        for i, chunk in enumerate(final_chunks):
            page_no = 1
            for key, pno in page_lookup.items():
                if key[:15] in chunk.page_content[:30] or chunk.page_content[:15] in key:
                    page_no = pno; break
            
            # 為非結構化文本注入 Nomic 檢索前綴
            chunk.page_content = f"search_document: {chunk.page_content}"
            
            print(f"     ... 正在為 PDF 內容進行 AI 打標 ({i+1}/{len(final_chunks)})")
            llm_tags = tagger.tag(chunk.page_content)
            md_headers = {k: v for k, v in chunk.metadata.items() if k in ["H1", "H2", "H3"]}

            chunk.metadata.update({
                "source": source_name, "type": "pdf_manual", "chunk_index": i, "page_number": page_no,
                "headers": json.dumps(md_headers, ensure_ascii=False),
                "ai_equipment": llm_tags.get("ai_equipment", "未知"),
                "ai_error_code": llm_tags.get("ai_error_code", "無"),
                "ai_category":  llm_tags.get("ai_category", "操作步驟")
            })
            docs.append(chunk)
            ids.append(generate_doc_id(source_name, i))
    except Exception as e:
        print(f"解析 PDF 失敗 ({file_path}): {e}")
    return docs, ids

def process_single_docx(file_path: str) -> Tuple[list, list]:
    docs, ids = [], []
    if not HAS_DOCX: return docs, ids
    try:
        doc = docx.Document(file_path)
        full_text = []
        source_name = os.path.basename(file_path)
        print(f"  -> 正在解析 Word: {source_name}")
        
        for para in doc.paragraphs:
            if not para.text.strip(): continue
            prefix = ""
            if para.style.name.startswith('Heading'):
                level = para.style.name[-1]
                if level.isdigit() and int(level) <= 6:
                    prefix = "#" * int(level) + " "
            full_text.append(prefix + para.text)

        md_text = "\n\n".join(full_text)
        md_splitter = MarkdownHeaderTextSplitter(headers_to_split_on=[("#", "H1"), ("##", "H2"), ("###", "H3")], strip_headers=False)
        text_splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200)
        final_chunks = text_splitter.split_documents(md_splitter.split_text(md_text))

        for i, chunk in enumerate(final_chunks):
            chunk.page_content = f"search_document: {chunk.page_content}"
            llm_tags = tagger.tag(chunk.page_content)
            md_headers = {k: v for k, v in chunk.metadata.items() if k in ["H1", "H2", "H3"]}
            chunk.metadata.update({
                "source": source_name, "type": "word_manual", "chunk_index": i,
                "headers": json.dumps(md_headers, ensure_ascii=False),
                "ai_equipment": llm_tags.get("ai_equipment", "未知"),
                "ai_error_code": llm_tags.get("ai_error_code", "無"),
                "ai_category":  llm_tags.get("ai_category", "操作步驟")
            })
            docs.append(chunk)
            ids.append(generate_doc_id(source_name, i))
    except Exception as e:
        print(f"解析 Word 失敗 ({file_path}): {e}")
    return docs, ids

def process_single_json(file_path: str) -> Tuple[list, list]:
    docs, ids = [], []
    try:
        source_name = os.path.basename(file_path)
        print(f"  -> 正在解析 JSON: {source_name}")
        with open(file_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
            
        if isinstance(data, list):
            texts = [json.dumps(item, ensure_ascii=False, indent=2) for item in data]
        else:
            texts = [json.dumps(data, ensure_ascii=False, indent=2)]

        text_splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200)
        
        for idx, t in enumerate(texts):
            chunks = text_splitter.split_text(t)
            for j, chunk_text in enumerate(chunks):
                global_index = f"{idx}_{j}"
                formatted_text = f"search_document: {chunk_text}"
                llm_tags = tagger.tag(formatted_text)
                metadata = {
                    "source": source_name, "type": "json_data", "chunk_index": global_index,
                    "ai_equipment": llm_tags.get("ai_equipment", "未知"),
                    "ai_error_code": llm_tags.get("ai_error_code", "無"),
                    "ai_category": llm_tags.get("ai_category", "其他")
                }
                docs.append(Document(page_content=formatted_text, metadata=metadata))
                ids.append(generate_doc_id(source_name, global_index))
    except Exception as e:
        print(f"讀取 JSON 失敗 ({file_path}): {e}")
    return docs, ids

def process_single_txt(file_path: str) -> Tuple[list, list]:
    docs, ids = [], []
    try:
        source_name = os.path.basename(file_path)
        print(f"  -> 正在解析 TXT: {source_name}")
        with open(file_path, 'r', encoding='utf-8') as f:
            text = f.read()

        text_splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200)
        chunks = text_splitter.split_text(text)

        for i, chunk_text in enumerate(chunks):
            formatted_text = f"search_document: {chunk_text}"
            llm_tags = tagger.tag(formatted_text)
            metadata = {
                "source": source_name, "type": "txt_log", "chunk_index": i,
                "ai_equipment": llm_tags.get("ai_equipment", "未知"),
                "ai_error_code": llm_tags.get("ai_error_code", "無"),
                "ai_category": llm_tags.get("ai_category", "其他")
            }
            docs.append(Document(page_content=formatted_text, metadata=metadata))
            ids.append(generate_doc_id(source_name, i))
    except Exception as e:
        print(f"讀取 TXT 失敗 ({file_path}): {e}")
    return docs, ids

# ==========================================================
# 🌟 4. 檔案處理路由與入庫邏輯
# ==========================================================
def run_preprocessing():
    print("啟動資料預處理程序 (支援 PDF, CSV, DOCX, JSON, TXT)...")
    embeddings = OllamaEmbeddings(model=EMBED_MODEL)
    all_docs, all_ids = [], []
    converter = get_advanced_pdf_converter()

    for file_path in glob(os.path.join(DATA_DIR, "*.*")):
        ext = os.path.splitext(file_path)[1].lower()
        if ext == ".csv": d, i = process_single_csv(file_path)
        elif ext == ".pdf": d, i = process_single_pdf(file_path, converter)
        elif ext in [".docx", ".doc"]: d, i = process_single_docx(file_path)
        elif ext == ".json": d, i = process_single_json(file_path)
        elif ext == ".txt": d, i = process_single_txt(file_path)
        else: continue
            
        all_docs.extend(d)
        all_ids.extend(i)

    if all_docs:
        unique_docs, unique_ids = [], []
        seen_ids = set()
        for doc, doc_id in zip(all_docs, all_ids):
            if doc_id not in seen_ids:
                seen_ids.add(doc_id)
                unique_docs.append(doc)
                unique_ids.append(doc_id)

        print(f"處理完成，準備寫入 {len(unique_docs)} 筆資料至 ChromaDB...")
        vector_store = Chroma(persist_directory=DB_DIR, embedding_function=embeddings)
        vector_store.add_documents(documents=unique_docs, ids=unique_ids)
        print(f"✅ 完成！資料庫位於: {DB_DIR}")
    else:
        print("無資料可處理。")

def save_upload_to_data_dir(temp_path: str) -> str:
    os.makedirs(DATA_DIR, exist_ok=True)
    dst = os.path.join(DATA_DIR, os.path.basename(temp_path))
    shutil.copy2(temp_path, dst)
    return dst

def ingest_uploaded_files_to_chroma(file_paths: List[str]) -> Dict:
    embeddings  = OllamaEmbeddings(model=EMBED_MODEL)
    vector_store = Chroma(persist_directory=DB_DIR, embedding_function=embeddings)
    all_docs, all_ids = [], []
    ok_files, bad_files = [], []
    converter = get_advanced_pdf_converter() 

    for p in file_paths:
        stable_path = save_upload_to_data_dir(p)
        ext = os.path.splitext(stable_path)[1].lower()
        
        if ext == ".csv": docs, ids = process_single_csv(stable_path)
        elif ext == ".pdf": docs, ids = process_single_pdf(stable_path, converter)
        elif ext in [".docx", ".doc"]: docs, ids = process_single_docx(stable_path)
        elif ext == ".json": docs, ids = process_single_json(stable_path)
        elif ext == ".txt": docs, ids = process_single_txt(stable_path)
        else:
            print(f"[入庫] 不支援的檔案類型，跳過: {stable_path}")
            docs, ids = [], []

        if docs:
            all_docs.extend(docs)
            all_ids.extend(ids)
            ok_files.append(os.path.basename(stable_path))
        else:
            bad_files.append(os.path.basename(stable_path))

    if all_docs:
        unique_docs, unique_ids = [], []
        seen_ids = set()
        for doc, doc_id in zip(all_docs, all_ids):
            if doc_id not in seen_ids:
                seen_ids.add(doc_id)
                unique_docs.append(doc)
                unique_ids.append(doc_id)
        vector_store.add_documents(documents=unique_docs, ids=unique_ids)

    return {"added_docs": len(all_docs), "ok_files": ok_files, "bad_files": bad_files, "db_dir": DB_DIR}

def delete_file_from_chroma(filename: str) -> bool:
    try:
        embeddings = OllamaEmbeddings(model=EMBED_MODEL)
        vector_store = Chroma(persist_directory=DB_DIR, embedding_function=embeddings)
        vector_store.delete(where={"source": filename})
        file_path = os.path.join(DATA_DIR, filename)
        if os.path.exists(file_path): os.remove(file_path)
        return True
    except Exception as e:
        return False

if __name__ == "__main__":
    run_preprocessing()